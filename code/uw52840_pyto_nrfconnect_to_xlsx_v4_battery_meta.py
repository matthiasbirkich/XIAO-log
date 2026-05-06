#!/usr/bin/env python3
# UW52840 nRF Connect CSV -> clean CSV/XLSX converter for Pyto and Thonny.
# v4: supports v22/v23/v24, including rhPct, vbat, and deployment metadata.
#
# This version fixes the main nRF Connect problem:
# BLE notifications are often split into 20-byte fragments such as
#   "idx,unix,tempC,bhMod"
#   "e,bhMtreg,bhRaw,lux,"
#   "ax,ay,az,wake,flags"
# and each measurement row may also arrive in several fragments.
#
# The script reconstructs the CSV from the HEX notification lines:
#   Updated Value of Characteristic A002 to ...
#
# Usage in Pyto:
#   1. Put this script and the nRF Connect export CSV in the same folder.
#   2. Run the script.
#   3. Choose the CSV file number.
#
# Usage in Thonny/PC:
#   python uw52840_pyto_nrfconnect_to_xlsx_v2.py UW52840-CLEAN.csv
#
# If openpyxl is unavailable, the script still writes a clean CSV.

INPUT_CSV = ""
OUTPUT_XLSX = ""
LOCAL_TZ = "Europe/Berlin"

import csv
import datetime as _dt
import re
import sys
from pathlib import Path

try:
    from zoneinfo import ZoneInfo
except Exception:
    ZoneInfo = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except Exception:
    Workbook = None


EXPECTED_HEADERS_V22 = [
    "idx", "unix", "tempC", "bhMode", "bhMtreg", "bhRaw",
    "lux", "ax", "ay", "az", "wake", "flags"
]

EXPECTED_HEADERS_V23 = [
    "idx", "unix", "tempC", "rhPct", "bhMode", "bhMtreg", "bhRaw",
    "lux", "ax", "ay", "az", "wake", "flags"
]

EXPECTED_HEADERS_V24 = [
    "idx", "unix", "tempC", "rhPct", "vbat", "bhMode", "bhMtreg", "bhRaw",
    "lux", "ax", "ay", "az", "wake", "flags"
]

# Default/fallback for current logger.
EXPECTED_HEADERS = EXPECTED_HEADERS_V24[:]

WAKE_MAP = {
    0: "unknown",
    1: "rtc",
    2: "reed",
    3: "powerup",
    4: "manual",
    99: "test"
}


def find_candidate_csv_files():
    seen = set()
    out = []
    places = []

    try:
        places.append(Path.cwd())
    except Exception:
        pass

    try:
        if "__file__" in globals():
            places.append(Path(__file__).resolve().parent)
    except Exception:
        pass

    for extra in [Path.home() / "Documents", Path.home() / "Downloads", Path.home()]:
        places.append(extra)

    for folder in places:
        try:
            if not folder.exists():
                continue
            for p in folder.glob("*.csv"):
                rp = str(p.resolve())
                if rp not in seen:
                    seen.add(rp)
                    out.append(p)
        except Exception:
            pass

    out.sort(key=lambda p: p.stat().st_mtime if p.exists() else 0, reverse=True)
    return out


def choose_input_path():
    if len(sys.argv) >= 2:
        return Path(sys.argv[1]).expanduser()

    if INPUT_CSV.strip():
        return Path(INPUT_CSV.strip()).expanduser()

    candidates = find_candidate_csv_files()
    if candidates:
        print("\nFound CSV files:")
        for i, p in enumerate(candidates[:30], start=1):
            print(f"  {i:2d}: {p}")

        choice = input("\nChoose a number, press Enter for 1, or type a full path: ").strip()
        if not choice:
            return candidates[0]
        if choice.isdigit():
            idx = int(choice)
            if 1 <= idx <= len(candidates[:30]):
                return candidates[idx - 1]
        return Path(choice).expanduser()

    return Path(input("Path to nRF Connect CSV export: ").strip()).expanduser()


def clean_base_name(input_path):
    stem = input_path.stem
    if stem.lower().endswith("_clean"):
        return stem
    return stem + "_clean"


def choose_output_path(input_path):
    if len(sys.argv) >= 3:
        return Path(sys.argv[2]).expanduser()

    if OUTPUT_XLSX.strip():
        return Path(OUTPUT_XLSX.strip()).expanduser()

    return input_path.with_name(clean_base_name(input_path) + ".xlsx")


def clean_csv_path(input_path, output_path):
    candidate = output_path.with_suffix(".csv")
    try:
        if candidate.resolve() == input_path.resolve():
            candidate = output_path.with_name(output_path.stem + "_converted.csv")
    except Exception:
        if str(candidate) == str(input_path):
            candidate = output_path.with_name(output_path.stem + "_converted.csv")
    return candidate


def decode_hex_notification_line(line):
    m = re.search(r"Updated Value of Characteristic A002 to ([0-9A-Fa-f ]+)\.", line)
    if not m:
        return None

    hx = m.group(1).replace(" ", "")
    if len(hx) < 2 or len(hx) % 2:
        return None

    try:
        return bytes.fromhex(hx).decode("utf-8", errors="replace")
    except Exception:
        return None


def extract_chunks_from_hex(text):
    chunks = []
    for line in text.splitlines():
        s = decode_hex_notification_line(line)
        if s is not None:
            chunks.append(s)
    return chunks


def extract_chunks_from_application_lines(text):
    chunks = []
    for line in text.splitlines():
        if "value received" not in line:
            continue

        if ",Application," in line:
            payload = line.split(",Application,", 1)[1]
        else:
            payload = line

        payload = payload.strip()
        payload = payload.replace('""', '"')

        marker = '" value received'
        pos = payload.find(marker)
        if pos >= 0:
            payload = payload[:pos]

        payload = payload.strip().strip('"')
        if payload:
            chunks.append(payload)

    return chunks


def choose_last_csv_block(chunks):
    blocks = []
    current = []
    in_block = False

    for chunk in chunks:
        if chunk == "CSV_BEGIN":
            current = []
            in_block = True
            continue

        if chunk == "CSV_END":
            if in_block:
                blocks.append(current)
            current = []
            in_block = False
            continue

        if in_block:
            current.append(chunk)

    if blocks:
        return blocks[-1]

    return chunks


def looks_like_record_start(fragment):
    return re.match(r"^\d+,\d{8,12},", fragment) is not None


def looks_like_metadata_start(fragment):
    return fragment.startswith("META,")


def looks_like_header_start(fragment):
    return fragment.startswith("idx,")


def reconstruct_csv_lines(chunks):
    block = choose_last_csv_block(chunks)

    lines = []
    current = None

    for frag in block:
        if frag in ("CSV_BEGIN", "CSV_END"):
            continue

        starts_new = (
            looks_like_metadata_start(frag)
            or looks_like_header_start(frag)
            or looks_like_record_start(frag)
        )

        if starts_new:
            if current is not None:
                lines.append(current)
            current = frag
        else:
            if current is None:
                current = frag
            else:
                current += frag

    if current is not None:
        lines.append(current)

    # Normalize old "R," prefix if present.
    cleaned = []
    for line in lines:
        line = line.strip()
        if line.startswith("R,"):
            line = line[2:]
        cleaned.append(line)

    # Ensure there is a header.
    if not any(line.startswith("idx,") for line in cleaned):
        cleaned.insert(0, ",".join(EXPECTED_HEADERS))

    return cleaned

def parse_reconstructed_lines(lines):
    if not lines:
        return [], [], {}

    metadata = {}
    data_lines = []

    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith("META,"):
            try:
                fields = next(csv.reader([line]))
            except Exception:
                fields = line.split(",")

            if len(fields) >= 3:
                key = fields[1].strip()
                value = ",".join(fields[2:]).strip()
                metadata[key] = value
            elif len(fields) == 2:
                metadata[fields[1].strip()] = ""
            continue

        data_lines.append(line)

    if not data_lines:
        return [], [], metadata

    try:
        header = next(csv.reader([data_lines[0]]))
    except Exception:
        header = data_lines[0].split(",")

    header = [h.strip().strip('"') for h in header]

    # Detect logger CSV format:
    # v22: idx,unix,tempC,bhMode,...
    # v23: idx,unix,tempC,rhPct,bhMode,...
    # v24: idx,unix,tempC,rhPct,vbat,bhMode,...
    if "vbat" in header:
        expected = EXPECTED_HEADERS_V24[:]
    elif "rhPct" in header:
        expected = EXPECTED_HEADERS_V23[:]
    else:
        expected = EXPECTED_HEADERS_V22[:]

    if header[: len(expected)] != expected:
        if set(expected).issubset(set(header)):
            pass
        else:
            expected = EXPECTED_HEADERS_V24[:]
            header = expected[:]

    rows = []
    for line in data_lines[1:]:
        if not line.strip():
            continue

        try:
            fields = next(csv.reader([line]))
        except Exception:
            fields = line.split(",")

        fields = [f.strip().strip('"') for f in fields]

        if len(fields) < len(expected):
            continue

        fields = fields[: len(expected)]
        item = dict(zip(expected, fields))

        # Normalize older files by adding missing current columns.
        if "rhPct" not in item:
            item["rhPct"] = ""
        if "vbat" not in item:
            item["vbat"] = ""

        rows.append(item)

    return header, rows, metadata

def to_int(value, default=None):
    try:
        if value is None or str(value).strip() == "":
            return default
        return int(float(str(value).strip()))
    except Exception:
        return default


def to_float(value, default=None):
    try:
        if value is None or str(value).strip() == "":
            return default
        return float(str(value).strip())
    except Exception:
        return default


def datetime_strings(unix_value):
    u = to_int(unix_value)
    if not u:
        return "", ""

    try:
        utc_dt = _dt.datetime.fromtimestamp(u, tz=_dt.timezone.utc)
    except Exception:
        return "", ""

    utc_s = utc_dt.strftime("%Y-%m-%d %H:%M:%S UTC")
    local_s = ""

    if ZoneInfo is not None:
        try:
            local_dt = utc_dt.astimezone(ZoneInfo(LOCAL_TZ))
            local_s = local_dt.strftime("%Y-%m-%d %H:%M:%S %Z")
        except Exception:
            pass

    return utc_s, local_s


def enrich_rows(rows, metadata=None):
    metadata = metadata or {}
    logger_id = metadata.get("logger_id", "")
    salinity_psu = to_float(metadata.get("salinity_psu"))
    lat_deg = to_float(metadata.get("lat_deg"))
    lon_deg = to_float(metadata.get("lon_deg"))

    enriched = []

    for r in rows:
        flags = to_int(r.get("flags"), 0) or 0
        wake_code = to_int(r.get("wake"), 0) or 0
        utc_s, local_s = datetime_strings(r.get("unix"))

        enriched.append({
            "idx": to_int(r.get("idx")),
            "unix": to_int(r.get("unix")),
            "utc_time": utc_s,
            "local_time": local_s,
            "logger_id": logger_id,
            "salinity_psu": salinity_psu,
            "lat_deg": lat_deg,
            "lon_deg": lon_deg,
            "tempC": to_float(r.get("tempC")),
            "rhPct": to_float(r.get("rhPct")),
            "vbat": to_float(r.get("vbat")),
            "lux": to_float(r.get("lux")),
            "ax_mg": to_int(r.get("ax")),
            "ay_mg": to_int(r.get("ay")),
            "az_mg": to_int(r.get("az")),
            "wake_code": wake_code,
            "wake_source": WAKE_MAP.get(wake_code, f"code_{wake_code}"),
            "flags": flags,
            "logging_enabled": bool(flags & 0x01),
            "rtc_valid": bool(flags & 0x02),
            "sht40_ok": bool(flags & 0x04),
            "bh1750_ok": bool(flags & 0x08),
            "imu_ok": bool(flags & 0x10),
            "bhMode": to_int(r.get("bhMode")),
            "bhMtreg": to_int(r.get("bhMtreg")),
            "bhRaw": to_int(r.get("bhRaw")),
        })

    return enriched


def write_clean_csv(path, rows):
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    fields = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def autosize(ws, max_width=34):
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        best = 10
        for cell in column_cells:
            if cell.value is None:
                continue
            best = max(best, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = best


def write_xlsx(path, input_path, reconstructed_lines, rows):
    if Workbook is None:
        return False

    wb = Workbook()

    ws = wb.active
    ws.title = "Measurements"

    headers = list(rows[0].keys()) if rows else [
        "idx", "unix", "utc_time", "local_time", "logger_id", "salinity_psu",
        "lat_deg", "lon_deg", "tempC", "rhPct", "vbat", "lux",
        "ax_mg", "ay_mg", "az_mg", "wake_code", "wake_source", "flags",
        "logging_enabled", "rtc_valid", "sht40_ok", "bh1750_ok", "imu_ok",
        "bhMode", "bhMtreg", "bhRaw"
    ]

    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    for col_name in ["tempC", "rhPct", "vbat", "lux", "salinity_psu", "lat_deg", "lon_deg"]:
        if col_name in headers:
            col = headers.index(col_name) + 1
            for column in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for cell in column:
                    cell.number_format = "0.00"

    autosize(ws)

    summary = wb.create_sheet("Summary")
    # Metadata is repeated per row for easy filtering, but also summarized here.
    first = rows[0] if rows else {}
    summary_rows = [
        ["Source file", str(input_path)],
        ["Records", len(rows)],
        ["Logger ID", first.get("logger_id", "")],
        ["Salinity PSU", first.get("salinity_psu", "")],
        ["Latitude deg", first.get("lat_deg", "")],
        ["Longitude deg", first.get("lon_deg", "")],
        ["Local timezone", LOCAL_TZ],
        ["Created", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]

    if rows:
        summary_rows.extend([
            ["First UTC", rows[0].get("utc_time", "")],
            ["Last UTC", rows[-1].get("utc_time", "")],
            ["First local", rows[0].get("local_time", "")],
            ["Last local", rows[-1].get("local_time", "")],
        ])

        temps = [r.get("tempC") for r in rows if isinstance(r.get("tempC"), (int, float))]
        rhs = [r.get("rhPct") for r in rows if isinstance(r.get("rhPct"), (int, float))]
        vbats = [r.get("vbat") for r in rows if isinstance(r.get("vbat"), (int, float))]
        luxs = [r.get("lux") for r in rows if isinstance(r.get("lux"), (int, float))]

        if temps:
            summary_rows.extend([
                ["Temperature avg °C", sum(temps) / len(temps)],
                ["Temperature min °C", min(temps)],
                ["Temperature max °C", max(temps)],
            ])

        if rhs:
            summary_rows.extend([
                ["Relative humidity avg %", sum(rhs) / len(rhs)],
                ["Relative humidity min %", min(rhs)],
                ["Relative humidity max %", max(rhs)],
            ])

        if vbats:
            summary_rows.extend([
                ["Battery avg V", sum(vbats) / len(vbats)],
                ["Battery min V", min(vbats)],
                ["Battery max V", max(vbats)],
            ])

        if luxs:
            summary_rows.extend([
                ["Lux avg", sum(luxs) / len(luxs)],
                ["Lux min", min(luxs)],
                ["Lux max", max(luxs)],
            ])

    for row in summary_rows:
        summary.append(row)

    for cell in summary["A"]:
        cell.font = Font(bold=True)
    autosize(summary, max_width=60)

    raw = wb.create_sheet("Reconstructed CSV")
    raw.append(["line"])
    for line in reconstructed_lines:
        raw.append([line])
    raw.freeze_panes = "A2"
    raw["A1"].font = Font(bold=True)
    raw["A1"].fill = header_fill
    raw.column_dimensions["A"].width = 90

    wb.save(path)
    return True


def convert_file(input_path, output_path):
    text = input_path.read_text(encoding="utf-8", errors="ignore")

    chunks = extract_chunks_from_hex(text)

    extraction_mode = "hex"
    if not chunks:
        chunks = extract_chunks_from_application_lines(text)
        extraction_mode = "application-text"

    reconstructed_lines = reconstruct_csv_lines(chunks)
    header, parsed_rows, metadata = parse_reconstructed_lines(reconstructed_lines)
    rows = enrich_rows(parsed_rows, metadata)

    clean_csv = clean_csv_path(input_path, output_path)
    write_clean_csv(clean_csv, rows)

    xlsx_ok = write_xlsx(output_path, input_path, reconstructed_lines, rows)

    return {
        "mode": extraction_mode,
        "chunks": len(chunks),
        "reconstructed_lines": len(reconstructed_lines),
        "records": len(rows),
        "clean_csv": clean_csv,
        "xlsx": output_path if xlsx_ok else None,
    }


def main():
    input_path = choose_input_path()

    if not input_path.exists():
        print(f"\nERROR: File not found: {input_path}")
        return 1

    output_path = choose_output_path(input_path)
    result = convert_file(input_path, output_path)

    print(f"\nExtraction mode: {result['mode']}")
    print(f"BLE notification chunks found: {result['chunks']}")
    print(f"Reconstructed CSV lines: {result['reconstructed_lines']}")
    print(f"Measurement records found: {result['records']}")
    print(f"Clean CSV written: {result['clean_csv']}")

    if result["xlsx"] is not None:
        print(f"XLSX written: {result['xlsx']}")
    else:
        print("\nXLSX was not created because openpyxl is not installed.")
        print("Pyto: install package 'openpyxl' and run again.")
        print("Thonny: Tools > Manage packages > openpyxl")

    if result["records"] == 0:
        print("\nWARNING: No records were reconstructed.")
        print("Make sure the nRF Connect export includes the CSV_BEGIN/CSV_END dump")
        print("and the 'Updated Value of Characteristic A002 to ...' lines.")

    print("\nDone.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
